from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from backpack_tax.app import create_app
from backpack_tax.db import TaxStore


def _insert_trade(
    store: TaxStore,
    *,
    occurred_at: str,
    asset: str,
    quantity: str,
    gross: str,
    source_type: str = "BACKPACK",
    fee_asset: str = "USDC",
) -> None:
    store.connection.execute(
        """
        INSERT INTO normalized_events (
            source_type, occurred_at_utc, event_type, asset, instrument_type,
            side, quantity, unit_price, gross_amount, fee_asset, fee_quantity,
            classification_status, created_at_utc
        ) VALUES (
            ?, ?, 'TRADE_FILL', ?, 'SPOT_CRYPTO', 'BUY', ?, '100',
            ?, ?, '0.1', 'NORMALIZED', '2026-01-01'
        )
        """,
        (source_type, occurred_at, asset, quantity, gross, fee_asset),
    )


def test_excel_export_uses_selected_period_and_numeric_cells(tmp_path: Path) -> None:
    app = create_app(project_root=tmp_path, testing=True)
    with TaxStore(tmp_path / "data" / "tax_ledger.db") as store:
        _insert_trade(
            store,
            occurred_at="2026-01-10T00:00:00+00:00",
            asset="JAN",
            quantity="2",
            gross="200",
        )
        _insert_trade(
            store,
            occurred_at="2026-02-10T00:00:00+00:00",
            asset="FEB",
            quantity="3",
            gross="300",
        )
        store.connection.commit()

    response = app.test_client().get(
        "/assets/export.xlsx?year=2026&month=1"
    )

    assert response.status_code == 200
    assert response.data[:2] == b"PK"
    assert "Backpack_PNL_2026-01.xlsx" in response.headers["Content-Disposition"]
    workbook = load_workbook(BytesIO(response.data), data_only=True)
    assert workbook.sheetnames == ["요약", "자산별 손익"]
    assert workbook["요약"]["B3"].value == "2026년 1월"
    detail = workbook["자산별 손익"]
    assert detail.auto_filter.ref == "A4:P5"
    assert detail["B5"].value == "JAN"
    assert detail["E5"].value == 200
    assert detail["F5"].value == 200
    assert detail["G5"].value == 0
    assert detail["F4"].value == "누적 매수액 (USD)"
    assert detail["G4"].value == "누적 매도액 (USD)"
    assert isinstance(detail["E5"].value, (int, float))


def test_excel_week_filename_uses_actual_month_end(tmp_path: Path) -> None:
    app = create_app(project_root=tmp_path, testing=True)

    response = app.test_client().get(
        "/assets/export.xlsx?year=2026&month=4&week=5"
    )

    assert response.status_code == 200
    assert (
        "Backpack_PNL_2026-04-29_to_2026-04-30.xlsx"
        in response.headers["Content-Disposition"]
    )


def test_excel_export_uses_selected_exchange_and_currency(tmp_path: Path) -> None:
    app = create_app(project_root=tmp_path, testing=True)
    with TaxStore(tmp_path / "data" / "tax_ledger.db") as store:
        _insert_trade(
            store,
            occurred_at="2026-01-10T00:00:00+00:00",
            asset="BTC",
            quantity="0.01",
            gross="1400000",
            source_type="UPBIT",
            fee_asset="KRW",
        )
        store.connection.commit()

    response = app.test_client().get(
        "/assets/export.xlsx?exchange=UPBIT&year=2026&month=1"
    )

    assert response.status_code == 200
    assert "Upbit_PNL_2026-01.xlsx" in response.headers["Content-Disposition"]
    workbook = load_workbook(BytesIO(response.data), data_only=True)
    assert workbook["요약"]["A1"].value == "업비트 자산별 손익 보고서"
    assert workbook["요약"]["B4"].value == "업비트"
    assert workbook["자산별 손익"]["F4"].value == "누적 매수액 (KRW)"
    assert workbook["자산별 손익"]["B5"].value == "BTC"
