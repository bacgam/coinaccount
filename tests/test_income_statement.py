from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from backpack_tax.analytics import asset_performance
from backpack_tax.app import create_app
from backpack_tax.db import TaxStore
from backpack_tax.income_statement import build_income_statement


def _insert_report_fixture(store: TaxStore) -> None:
    store.connection.executemany(
        """
        INSERT INTO normalized_events (
            source_type, occurred_at_utc, event_type, asset, instrument_type,
            side, quantity, unit_price, gross_amount, fee_asset, fee_quantity,
            classification_status, created_at_utc
        ) VALUES (
            'BACKPACK', ?, 'TRADE_FILL', 'TEST', 'SPOT_CRYPTO', ?, ?, ?, ?,
            'USDC', ?, 'NORMALIZED', '2026-01-01'
        )
        """,
        [
            ("2025-12-31T00:00:00+00:00", "BUY", "2", "100", "200", "0"),
            ("2026-01-10T00:00:00+00:00", "SELL", "1", "150", "150", "2"),
        ],
    )
    store.connection.execute(
        """
        INSERT INTO normalized_events (
            source_type, occurred_at_utc, event_type, asset, instrument_type,
            realized_pnl, funding_quantity, interest_quantity,
            trading_fee_quantity, liquidation_fee_quantity,
            classification_status, created_at_utc
        ) VALUES (
            'BACKPACK', '2026-01-11T00:00:00+00:00',
            'FUTURES_POSITION_CLOSED', 'BTC', 'PERPETUAL_FUTURE',
            '20', '-1', '0.5', '2', '0.1', 'NORMALIZED', '2026-01-11'
        )
        """
    )
    store.connection.commit()


def test_income_statement_uses_prior_cost_basis_and_separates_futures(
    tmp_path: Path,
) -> None:
    with TaxStore(tmp_path / "tax.db") as store:
        _insert_report_fixture(store)
        rows = asset_performance(
            store,
            start_utc=datetime(2026, 1, 1, tzinfo=UTC),
            end_utc=datetime(2026, 2, 1, tzinfo=UTC),
            source_type="BACKPACK",
        )
    statement = build_income_statement(rows)

    assert statement["recognized_sales"] == Decimal("150")
    assert statement["cost_of_sales"] == Decimal("100")
    assert statement["spot_profit"] == Decimal("50")
    assert statement["fees"] == Decimal("4.1")
    assert statement["future_income"] == Decimal("18.5")
    assert statement["future_expense"] == Decimal("0")
    assert statement["net_profit"] == Decimal("64.4")
    assert statement["lines"][1]["amount_display"] == "(100)"
    assert statement["lines"][-1]["amount_display"] == "64.4"


def test_report_preview_and_excel_use_selected_period(tmp_path: Path) -> None:
    app = create_app(project_root=tmp_path, testing=True)
    with TaxStore(tmp_path / "data" / "tax_ledger.db") as store:
        _insert_report_fixture(store)
    client = app.test_client()

    page = client.get("/reports?exchange=BACKPACK&year=2026&month=1")
    assert page.status_code == 200
    assert "보고서 추출".encode() in page.data
    assert "손 익 계 산 서".encode() in page.data
    assert "2026년 1월".encode() in page.data
    assert "현물 매도금액(취득가 확인분)".encode() in page.data
    assert "기간 순손익".encode() in page.data
    assert b"64.4" in page.data
    assert "확인한 내용 Excel 추출".encode() in page.data

    response = client.get(
        "/reports/export.xlsx?exchange=BACKPACK&year=2026&month=1"
    )
    assert response.status_code == 200
    assert response.data[:2] == b"PK"
    assert (
        "Backpack_Income_Statement_2026-01.xlsx"
        in response.headers["Content-Disposition"]
    )
    workbook = load_workbook(BytesIO(response.data), data_only=True)
    assert workbook.sheetnames == ["손익계산서", "산출근거"]
    statement_sheet = workbook["손익계산서"]
    assert statement_sheet["A1"].value == "손 익 계 산 서"
    assert statement_sheet["A13"].value == "X."
    assert statement_sheet["D13"].value == 64.4
    assert "(" in statement_sheet["D5"].number_format
    assert workbook["산출근거"]["A2"].value in {"BTC", "TEST"}


def test_report_uses_krw_for_korean_exchanges(tmp_path: Path) -> None:
    app = create_app(project_root=tmp_path, testing=True)
    page = app.test_client().get("/reports?exchange=BITHUMB")

    assert page.status_code == 200
    assert "빗썸 · 전체 기간".encode() in page.data
    assert "단위: KRW".encode() in page.data
