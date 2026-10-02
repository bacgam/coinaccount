from __future__ import annotations

import io
from datetime import datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps

from backpack_tax.periods import PeriodSelection, SEOUL


GREEN = "1E6B4F"
LIGHT_GREEN = "DFF2AA"
DARK = "17211B"
MUTED = "69736D"
LIGHT_LINE = "DFE5E1"
RED = "B53E36"
WHITE = "FFFFFF"


def _safe_text(value: object) -> str:
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _number(value: object) -> float:
    return float(Decimal(str(value or "0")))


def _period_boundaries(period: PeriodSelection) -> str:
    if period.start_utc is None or period.end_utc is None:
        return "전체 보유 자료"
    start = period.start_utc.astimezone(SEOUL)
    end = period.end_utc.astimezone(SEOUL)
    return f"{start:%Y-%m-%d %H:%M} 이상 / {end:%Y-%m-%d %H:%M} 미만 (KST)"


def _period_suffix(period: PeriodSelection) -> str:
    if period.period_type == "all":
        return "all"
    elif period.period_type == "year":
        return period.year_value
    elif period.period_type == "month":
        return f"{int(period.year_value):04d}-{int(period.month_value):02d}"
    elif period.period_type == "week":
        start_day = period.start_utc.astimezone(SEOUL).date()
        final_day = period.end_utc.astimezone(SEOUL).date() - timedelta(days=1)
        return f"{start_day.isoformat()}_to_{final_day.isoformat()}"
    return (
        f"{int(period.year_value):04d}-{int(period.month_value):02d}"
        f"-{int(period.day_value):02d}"
    )


def excel_filename(period: PeriodSelection, source_type: str = "BACKPACK") -> str:
    suffix = _period_suffix(period)
    exchange_prefix = {
        "BACKPACK": "Backpack",
        "UPBIT": "Upbit",
        "BITHUMB": "Bithumb",
    }.get(source_type.upper(), "Backpack")
    return f"{exchange_prefix}_PNL_{suffix}.xlsx"


def income_statement_filename(
    period: PeriodSelection, source_type: str = "BACKPACK"
) -> str:
    exchange_prefix = {
        "BACKPACK": "Backpack",
        "UPBIT": "Upbit",
        "BITHUMB": "Bithumb",
    }.get(source_type.upper(), "Backpack")
    return f"{exchange_prefix}_Income_Statement_{_period_suffix(period)}.xlsx"


def _style_title(cell) -> None:
    cell.font = Font(name="맑은 고딕", size=18, bold=True, color=DARK)
    cell.alignment = Alignment(vertical="center")


def _style_header_row(worksheet, row: int, start_column: int, end_column: int) -> None:
    thin = Side(style="thin", color=LIGHT_LINE)
    for column in range(start_column, end_column + 1):
        cell = worksheet.cell(row=row, column=column)
        cell.font = Font(name="맑은 고딕", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)


def _logo_image(path: Path, buffers: list[io.BytesIO]) -> ExcelImage | None:
    try:
        with Image.open(path) as source:
            source.load()
            image = source.convert("RGBA")
    except (OSError, ValueError):
        return None
    image = ImageOps.contain(image, (40, 40), Image.Resampling.LANCZOS)
    canvas = Image.new("RGBA", (48, 48), (255, 255, 255, 0))
    canvas.alpha_composite(image, ((48 - image.width) // 2, (48 - image.height) // 2))
    buffer = io.BytesIO()
    canvas.save(buffer, format="PNG", optimize=True)
    buffer.seek(0)
    buffers.append(buffer)
    logo = ExcelImage(buffer)
    logo.width = 28
    logo.height = 28
    return logo


def build_asset_performance_workbook(
    *,
    period: PeriodSelection,
    rows: list[dict[str, Any]],
    summary: dict[str, str | int],
    logo_directory: Path,
    exchange_name: str = "Backpack",
    currency_code: str = "USD",
) -> io.BytesIO:
    workbook = Workbook()
    workbook.properties.creator = "코인어카운트"
    workbook.properties.title = f"{period.label} 자산별 손익"
    workbook.properties.description = "한국시간 기준 자산별 거래량, 실현손익 및 수수료"
    created_at = datetime.now(SEOUL)

    overview = workbook.active
    overview.title = "요약"
    overview.sheet_view.showGridLines = False
    overview.merge_cells("A1:H1")
    overview["A1"] = f"{exchange_name} 자산별 손익 보고서"
    _style_title(overview["A1"])
    overview.row_dimensions[1].height = 30
    metadata = (
        ("선택 기간", period.label),
        ("거래소", exchange_name),
        ("기간 경계", _period_boundaries(period)),
        ("시간대", "Asia/Seoul (KST)"),
        ("생성 시각", created_at.strftime("%Y-%m-%d %H:%M:%S KST")),
    )
    for row_number, (label, value) in enumerate(metadata, start=3):
        overview.cell(row_number, 1, label).font = Font(bold=True, color=GREEN)
        overview.cell(row_number, 2, _safe_text(value))
    overview.append([])
    overview.append(["구분", "값", "단위/설명"])
    summary_row = overview.max_row
    _style_header_row(overview, summary_row, 1, 3)
    summary_values = (
        ("거래 자산", int(summary["asset_count"]), "개"),
        ("거래량", _number(str(summary["volume"]).replace(",", "")), currency_code),
        ("누적 PNL", _number(str(summary["pnl"]).replace(",", "")), f"{currency_code} · 수수료 제외"),
        ("수수료", _number(str(summary["fees"]).replace(",", "")), f"{currency_code} · 별도 합계"),
    )
    for item in summary_values:
        overview.append(item)
    for row_number in range(summary_row + 1, summary_row + 5):
        overview.cell(row_number, 2).number_format = "#,##0.00;[Red]-#,##0.00"
    note_row = overview.max_row + 3
    overview.cell(note_row, 1, "계산 기준").font = Font(bold=True, color=GREEN)
    overview.cell(
        note_row + 1,
        1,
        "현물·토큰화 주식은 이동평균 취득가 기준 실현손익입니다.",
    )
    overview.cell(
        note_row + 2,
        1,
        "무기한 선물은 종료 포지션 실현손익과 펀딩을 반영합니다.",
    )
    overview.cell(
        note_row + 3,
        1,
        "수수료는 PNL에서 차감하지 않고 별도 열로 표시합니다.",
    )
    overview.column_dimensions["A"].width = 24
    overview.column_dimensions["B"].width = 42
    overview.column_dimensions["C"].width = 22

    detail = workbook.create_sheet("자산별 손익")
    detail.sheet_view.showGridLines = False
    detail.merge_cells("A1:P1")
    detail["A1"] = f"{exchange_name} · {period.label} · 자산별 손익"
    _style_title(detail["A1"])
    detail["A2"] = _period_boundaries(period)
    detail["A2"].font = Font(name="맑은 고딕", size=9, color=MUTED)
    headers = (
        "로고",
        "자산코드",
        "이름",
        "상품유형",
        f"거래량 ({currency_code})",
        f"누적 매수액 ({currency_code})",
        f"누적 매도액 ({currency_code})",
        f"누적 PNL ({currency_code}, 수수료 제외)",
        f"수수료 ({currency_code})",
        "취득가 미확인 매도량",
        "상태",
        "내부 이동 입금량",
        "내부 이동 출금량",
        "네트워크 수수료 수량",
        f"네트워크 수수료 원가 ({currency_code})",
        "가격 증빙 대기 수량",
    )
    for column, header in enumerate(headers, start=1):
        detail.cell(4, column, header)
    _style_header_row(detail, 4, 1, len(headers))
    detail.row_dimensions[4].height = 24
    image_buffers: list[io.BytesIO] = []
    logo_root = logo_directory.resolve()
    for row_number, item in enumerate(rows, start=5):
        values = (
            "",
            _safe_text(item["asset"]),
            _safe_text(item["name"]),
            _safe_text(item["instrument_label"]),
            _number(item["volume"]),
            _number(item["buy_amount"]),
            _number(item["sell_amount"]),
            _number(item["pnl"]),
            _number(item["fees"]),
            _number(item["unknown_sell_quantity"]),
            "가격 증빙 필요"
            if item["has_pending_transfer_basis"]
            else "취득가 확인 필요"
            if item["has_unknown_basis"]
            else "계산 완료",
            _number(item["transfer_in_quantity"]),
            _number(item["transfer_out_quantity"]),
            _number(item["transfer_fee_quantity"]),
            _number(item["transfer_fee_cost"]),
            _number(item["pending_transfer_basis_quantity"]),
        )
        for column, value in enumerate(values, start=1):
            detail.cell(row_number, column, value)
        detail.row_dimensions[row_number].height = 28
        logo_filename = item.get("logo_filename")
        if logo_filename:
            logo_path = (logo_root / Path(str(logo_filename)).name).resolve()
            if logo_path.is_relative_to(logo_root) and logo_path.is_file():
                logo = _logo_image(logo_path, image_buffers)
                if logo:
                    detail.add_image(logo, f"A{row_number}")
        pnl_cell = detail.cell(row_number, 8)
        pnl_cell.font = Font(color=GREEN if item["pnl"] >= 0 else RED, bold=True)
        if item["has_unknown_basis"]:
            detail.cell(row_number, 11).fill = PatternFill("solid", fgColor="FFF1D6")

    last_row = max(4, detail.max_row)
    detail.auto_filter.ref = f"A4:P{last_row}"
    detail.freeze_panes = "A5"
    detail.sheet_properties.pageSetUpPr.fitToPage = True
    detail.page_setup.orientation = "landscape"
    detail.page_setup.fitToWidth = 1
    currency_symbol = "$" if currency_code == "USD" else "₩"
    for column in range(5, 11):
        number_format = (
            f"{currency_symbol}#,##0.00;[Red]-{currency_symbol}#,##0.00"
            if column in {5, 6, 7, 8, 9}
            else "#,##0.########"
        )
        for row_number in range(5, last_row + 1):
            detail.cell(row_number, column).number_format = number_format
    for column in range(12, 17):
        for row_number in range(5, last_row + 1):
            detail.cell(row_number, column).number_format = (
                f"{currency_symbol}#,##0.00;[Red]-{currency_symbol}#,##0.00"
                if column == 15
                else "#,##0.########"
            )
    widths = (9, 15, 24, 18, 22, 18, 18, 24, 16, 24, 18, 20, 20, 22, 24, 22)
    for column, width in enumerate(widths, start=1):
        detail.column_dimensions[get_column_letter(column)].width = width
    for row in detail.iter_rows(min_row=4, max_row=last_row, min_col=1, max_col=16):
        for cell in row:
            cell.alignment = Alignment(vertical="center")

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_income_statement_workbook(
    *,
    period: PeriodSelection,
    exchange_name: str,
    currency_code: str,
    statement: dict[str, Any],
    asset_rows: list[dict[str, Any]],
) -> io.BytesIO:
    workbook = Workbook()
    workbook.properties.creator = "코인어카운트"
    workbook.properties.title = f"{exchange_name} {period.label} 관리용 손익계산서"
    workbook.properties.description = "거래소 원장 기반 기간별 관리용 손익계산서"

    sheet = workbook.active
    sheet.title = "손익계산서"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:D1")
    sheet["A1"] = "손 익 계 산 서"
    sheet["A1"].font = Font(name="맑은 고딕", size=20, bold=True, color=DARK)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:D2")
    sheet["A2"] = f"{exchange_name} · {period.label} · 단위: {currency_code}"
    sheet["A2"].font = Font(name="맑은 고딕", size=11, bold=True, color=DARK)
    sheet["A2"].alignment = Alignment(horizontal="center")
    sheet.row_dimensions[1].height = 34
    sheet.row_dimensions[2].height = 23

    medium = Side(style="medium", color=DARK)
    thin = Side(style="thin", color="707070")
    start_row = 4
    for offset, line in enumerate(statement["lines"]):
        row = start_row + offset
        sheet.cell(row, 1, line["roman"])
        sheet.cell(row, 2, line["label"])
        sheet.cell(row, 4, _number(line["amount"]))
        sheet.cell(row, 1).font = Font(name="맑은 고딕", size=11, bold=True)
        sheet.cell(row, 2).font = Font(name="맑은 고딕", size=11, bold=True)
        sheet.cell(row, 4).font = Font(name="맑은 고딕", size=11, bold=True)
        sheet.cell(row, 4).alignment = Alignment(horizontal="right")
        sheet.cell(row, 4).number_format = "#,##0.00;(#,##0.00);0"
        sheet.row_dimensions[row].height = 23
        border_side = medium if offset == 0 else thin if line["subtotal"] else None
        if border_side:
            for column in range(1, 5):
                sheet.cell(row, column).border = Border(top=border_side)
        if line["final"]:
            for column in (1, 2, 4):
                sheet.cell(row, column).font = Font(
                    name="맑은 고딕", size=12, bold=True, color="0000FF"
                )
            for column in range(1, 5):
                sheet.cell(row, column).border = Border(top=medium, bottom=medium)

    note_row = start_row + len(statement["lines"]) + 2
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=4)
    sheet.cell(note_row, 1, "※ 세금은 자동 계산하지 않으며 신고용 확정 전 전문가 검토가 필요합니다.")
    sheet.cell(note_row, 1).font = Font(name="맑은 고딕", size=9, color=MUTED)
    if statement["has_unknown_basis"]:
        sheet.merge_cells(
            start_row=note_row + 1,
            start_column=1,
            end_row=note_row + 1,
            end_column=4,
        )
        sheet.cell(
            note_row + 1,
            1,
            f"※ 취득가 미확인 매도금액 {statement['unknown_basis_display']} {currency_code}는 실현손익에서 제외했습니다.",
        )
        sheet.cell(note_row + 1, 1).font = Font(
            name="맑은 고딕", size=9, bold=True, color=RED
        )
    if statement.get("has_pending_transfer_basis"):
        pending_text = ", ".join(
            f"{item['asset']} {item['quantity']}"
            for item in statement["pending_transfer_basis"]
        )
        sheet.merge_cells(
            start_row=note_row + 2,
            start_column=1,
            end_row=note_row + 2,
            end_column=4,
        )
        sheet.cell(
            note_row + 2,
            1,
            f"※ USD/USDT↔KRW 가격 증빙 대기 내부 이동: {pending_text}",
        )
        sheet.cell(note_row + 2, 1).font = Font(
            name="맑은 고딕", size=9, bold=True, color=RED
        )
    sheet.column_dimensions["A"].width = 9
    sheet.column_dimensions["B"].width = 38
    sheet.column_dimensions["C"].width = 5
    sheet.column_dimensions["D"].width = 22
    sheet.freeze_panes = "A4"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.orientation = "portrait"

    detail = workbook.create_sheet("산출근거")
    detail.sheet_view.showGridLines = False
    detail.append(
        [
            "자산",
            "상품유형",
            f"거래량 ({currency_code})",
            f"매수액 ({currency_code})",
            f"매도액 ({currency_code})",
            f"취득가 확인 매도액 ({currency_code})",
            f"처분원가 ({currency_code})",
            f"실현손익 ({currency_code})",
            f"수수료 ({currency_code})",
            f"취득가 미확인 매도액 ({currency_code})",
            "내부 이동 입금량",
            "내부 이동 출금량",
            f"네트워크 수수료 원가 ({currency_code})",
            "가격 증빙 대기 수량",
        ]
    )
    _style_header_row(detail, 1, 1, 14)
    for item in asset_rows:
        detail.append(
            [
                _safe_text(item["asset"]),
                _safe_text(item["instrument_label"]),
                _number(item["volume"]),
                _number(item["buy_amount"]),
                _number(item["sell_amount"]),
                _number(item.get("matched_sell_amount", 0)),
                _number(item.get("realized_cost", 0)),
                _number(item["pnl"]),
                _number(item["fees"]),
                _number(item.get("unknown_sell_amount", 0)),
                _number(item.get("transfer_in_quantity", 0)),
                _number(item.get("transfer_out_quantity", 0)),
                _number(item.get("transfer_fee_cost", 0)),
                _number(item.get("pending_transfer_basis_quantity", 0)),
            ]
        )
    last_row = max(1, detail.max_row)
    detail.auto_filter.ref = f"A1:N{last_row}"
    detail.freeze_panes = "A2"
    for row in detail.iter_rows(min_row=2, max_row=last_row, min_col=3, max_col=14):
        for cell in row:
            cell.number_format = "#,##0.00;(#,##0.00);0"
    for column, width in enumerate((15, 20, 18, 18, 18, 24, 18, 18, 18, 28, 20, 20, 24, 24), 1):
        detail.column_dimensions[get_column_letter(column)].width = width

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
